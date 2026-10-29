import random
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
from openpyxl.chart import PieChart, Reference
from openpyxl.chart.label import DataLabelList

VERDE = "54806A"

# 1. Generar 25 estudiantes ficticios
def generar_datos():
    estudiantes = []
    for i in range(1, 26):
        estudiantes.append({
            "id": i,
            "intentos": random.randint(5, 30),
            "resueltos": random.randint(1, 15),
            "puntaje": random.randint(100, 1000)
        })
    return estudiantes


# 2. Merge Sort genérico: sirve para puntajes y aristas
def combinar(izquierda, derecha, clave):
    resultado = []
    i = j = 0
    while i < len(izquierda) and j < len(derecha):
        if clave(izquierda[i]) <= clave(derecha[j]):
            resultado.append(izquierda[i])
            i += 1
        else:
            resultado.append(derecha[j])
            j += 1
    resultado.extend(izquierda[i:])
    resultado.extend(derecha[j:])
    return resultado


def merge_sort(datos, clave):
    if len(datos) <= 1:
        return datos
    mitad = len(datos) // 2
    izquierda = merge_sort(datos[:mitad], clave)
    derecha = merge_sort(datos[mitad:], clave)
    return combinar(izquierda, derecha, clave)


# 3. Voraz: escoger los 10 estudiantes con mayor puntaje
def seleccionar_voraz(estudiantes, cupos):
    ordenados = merge_sort(estudiantes.copy(), lambda e: -e["puntaje"])
    return ordenados[:cupos]


# 4. Programación dinámica: mochila 0/1.
# Cada estudiante cuesta su cantidad de intentos y aporta su puntaje.
def seleccionar_dinamica(estudiantes, presupuesto):
    n = len(estudiantes)
    dp = [[0] * (presupuesto + 1) for _ in range(n + 1)]

    for i in range(1, n + 1):
        estudiante = estudiantes[i - 1]
        costo = estudiante["intentos"]
        valor = estudiante["puntaje"]
        for capacidad in range(presupuesto + 1):
            dp[i][capacidad] = dp[i - 1][capacidad]
            if costo <= capacidad:
                alternativa = dp[i - 1][capacidad - costo] + valor
                if alternativa > dp[i][capacidad]:
                    dp[i][capacidad] = alternativa

    seleccionados = []
    capacidad = presupuesto
    for i in range(n, 0, -1):
        if dp[i][capacidad] != dp[i - 1][capacidad]:
            estudiante = estudiantes[i - 1]
            seleccionados.append(estudiante)
            capacidad -= estudiante["intentos"]
    seleccionados.reverse()
    return seleccionados, dp[n][presupuesto]


# 5. Kruskal: cada estudiante es un nodo y cada arista pesa
# la diferencia absoluta entre los puntajes de dos estudiantes.
def crear_aristas(estudiantes):
    aristas = []
    for i in range(len(estudiantes)):
        for j in range(i + 1, len(estudiantes)):
            a = estudiantes[i]
            b = estudiantes[j]
            aristas.append({
                "id_a": a["id"],
                "id_b": b["id"],
                "puntaje_a": a["puntaje"],
                "puntaje_b": b["puntaje"],
                "distancia": abs(a["puntaje"] - b["puntaje"])
            })
    return aristas


def buscar_padre(padres, nodo):
    while padres[nodo] != nodo:
        padres[nodo] = padres[padres[nodo]]
        nodo = padres[nodo]
    return nodo


def kruskal(estudiantes):
    aristas = crear_aristas(estudiantes)
    # Se reutiliza Merge Sort para ordenar aristas por distancia.
    ordenadas = merge_sort(aristas.copy(), lambda arista: arista["distancia"])
    padres = {e["id"]: e["id"] for e in estudiantes}
    arbol = []

    for arista in ordenadas:
        padre_a = buscar_padre(padres, arista["id_a"])
        padre_b = buscar_padre(padres, arista["id_b"])
        if padre_a != padre_b:
            padres[padre_a] = padre_b
            arbol.append(arista)
        if len(arbol) == len(estudiantes) - 1:
            break
    return arbol, ordenadas


# 6. Funciones para crear tablas en Excel
def dar_formato(hoja):
    for celda in hoja[1]:
        celda.font = Font(bold=True, color="FFFFFF")
        celda.fill = PatternFill(fill_type="solid", fgColor=VERDE)
        celda.alignment = Alignment(horizontal="center")
    for columna in hoja.columns:
        letra = get_column_letter(columna[0].column)
        ancho = max(len(str(c.value or "")) for c in columna)
        hoja.column_dimensions[letra].width = min(ancho + 3, 30)
    for fila in hoja.iter_rows(min_row=2):
        for celda in fila:
            celda.alignment = Alignment(horizontal="center")
    hoja.freeze_panes = "A2"
    if hoja.max_row > 1:
        hoja.auto_filter.ref = hoja.dimensions


def crear_hoja_estudiantes(libro, nombre, estudiantes):
    hoja = libro.create_sheet(nombre)
    hoja.append(["ID", "Intentos", "Resueltos", "Puntaje"])
    for e in estudiantes:
        hoja.append([e["id"], e["intentos"], e["resueltos"], e["puntaje"]])
    dar_formato(hoja)
    return hoja


def crear_hoja_resultados(libro, nombre, encabezados, filas):
    hoja = libro.create_sheet(nombre)
    hoja.append(encabezados)
    for fila in filas:
        hoja.append(fila)
    dar_formato(hoja)
    return hoja


# 7. Gráficas de torta simples, con porcentajes

def crear_hoja_graficas(libro, estudiantes):
    datos = libro.create_sheet("Datos gráficas")
    datos.append(["Rango de puntaje", "Cantidad"])
    rangos_puntaje = [("100-299", 100, 299), ("300-499", 300, 499),
                      ("500-699", 500, 699), ("700-899", 700, 899),
                      ("900-1000", 900, 1000)]
    for nombre, minimo, maximo in rangos_puntaje:
        datos.append([nombre, sum(1 for e in estudiantes if minimo <= e["puntaje"] <= maximo)])

    fila_intentos = datos.max_row + 3
    datos.cell(fila_intentos, 1, "Rango de intentos")
    datos.cell(fila_intentos, 2, "Cantidad")
    rangos_intentos = [("5-10", 5, 10), ("11-15", 11, 15), ("16-20", 16, 20),
                       ("21-25", 21, 25), ("26-30", 26, 30)]
    for fila, (nombre, minimo, maximo) in enumerate(rangos_intentos, fila_intentos + 1):
        datos.cell(fila, 1, nombre)
        datos.cell(fila, 2, sum(1 for e in estudiantes if minimo <= e["intentos"] <= maximo))

    fila_resueltos = datos.max_row + 3
    datos.cell(fila_resueltos, 1, "Rango de ejercicios resueltos")
    datos.cell(fila_resueltos, 2, "Cantidad")
    rangos_resueltos = [("1-3", 1, 3), ("4-6", 4, 6), ("7-9", 7, 9),
                        ("10-12", 10, 12), ("13-15", 13, 15)]
    for fila, (nombre, minimo, maximo) in enumerate(rangos_resueltos, fila_resueltos + 1):
        datos.cell(fila, 1, nombre)
        datos.cell(fila, 2, sum(1 for e in estudiantes if minimo <= e["resueltos"] <= maximo))

    for fila in (1, fila_intentos, fila_resueltos):
        for celda in datos[fila]:
            if celda.value is not None:
                celda.font = Font(bold=True, color="FFFFFF")
                celda.fill = PatternFill(fill_type="solid", fgColor=VERDE)

    graficas = libro.create_sheet("Gráficas")
    graficas.sheet_view.showGridLines = False

    def agregar_torta(titulo, inicio, fin, ancla):
        grafica = PieChart()
        grafica.title = titulo
        grafica.height = 7
        grafica.width = 11
        valores = Reference(datos, min_col=2, min_row=inicio, max_row=fin)
        categorias = Reference(datos, min_col=1, min_row=inicio, max_row=fin)
        grafica.add_data(valores)
        grafica.set_categories(categorias)
        grafica.dataLabels = DataLabelList()
        grafica.dataLabels.showPercent = True
        grafica.dataLabels.showLeaderLines = True
        grafica.legend.position = "r"
        graficas.add_chart(grafica, ancla)

    agregar_torta("Distribución de puntajes", 2, 6, "A1")
    agregar_torta("Distribución de intentos", fila_intentos + 1, fila_intentos + 5, "I1")
    agregar_torta("Ejercicios resueltos", fila_resueltos + 1, fila_resueltos + 5, "A16")
    datos.sheet_state = "hidden"


# 8. Programa principal
def main():
    random.seed(42)
    estudiantes = generar_datos()
    ordenados = merge_sort(estudiantes.copy(), lambda e: e["puntaje"])

    # Voraz: elige los 10 estudiantes con puntaje más alto.
    cupos = 10
    seleccion_voraz = seleccionar_voraz(estudiantes, cupos)

    # DP: maximiza puntaje sin superar 100 intentos sumados.
    presupuesto = 100
    seleccion_dp, puntaje_dp = seleccionar_dinamica(estudiantes, presupuesto)
    puntaje_voraz = sum(e["puntaje"] for e in seleccion_voraz)
    intentos_voraz = sum(e["intentos"] for e in seleccion_voraz)
    intentos_dp = sum(e["intentos"] for e in seleccion_dp)

    arbol, aristas_ordenadas = kruskal(estudiantes)

    libro = Workbook()
    libro.remove(libro.active)
    crear_hoja_estudiantes(libro, "Datos originales", estudiantes)
    crear_hoja_estudiantes(libro, "Ordenados por puntaje", ordenados)

    crear_hoja_resultados(libro, "Selección voraz",
        ["ID", "Intentos", "Resueltos", "Puntaje"],
        [[e["id"], e["intentos"], e["resueltos"], e["puntaje"]] for e in seleccion_voraz])

    crear_hoja_resultados(libro, "Programación dinámica",
        ["ID", "Intentos (costo)", "Resueltos", "Puntaje (valor)"],
        [[e["id"], e["intentos"], e["resueltos"], e["puntaje"]] for e in seleccion_dp])

    crear_hoja_resultados(libro, "Kruskal árbol mínimo",
        ["ID estudiante A", "ID estudiante B", "Puntaje A", "Puntaje B", "Diferencia"],
        [[a["id_a"], a["id_b"], a["puntaje_a"], a["puntaje_b"], a["distancia"]] for a in arbol])

    crear_hoja_resultados(libro, "Aristas ordenadas",
        ["ID estudiante A", "ID estudiante B", "Puntaje A", "Puntaje B", "Diferencia"],
        [[a["id_a"], a["id_b"], a["puntaje_a"], a["puntaje_b"], a["distancia"]] for a in aristas_ordenadas])

    crear_hoja_graficas(libro, estudiantes)
    nombre_archivo = "resultados_estudiantes.xlsx"
    libro.save(nombre_archivo)

    print("Excel creado:", nombre_archivo)
    print("Estudiantes simulados:", len(estudiantes))
    print("Voraz: cupos =", cupos, "| puntaje total =", puntaje_voraz,
          "| intentos utilizados =", intentos_voraz)
    print("Programación dinámica: presupuesto =", presupuesto,
          "| puntaje óptimo =", puntaje_dp, "| intentos utilizados =", intentos_dp)
    print("Kruskal: aristas del árbol mínimo =", len(arbol),
          "| suma de diferencias =", sum(a["distancia"] for a in arbol))
    print("Merge Sort: complejidad temporal O(n log n); espacial O(n).")


if __name__ == "__main__":
    main()
